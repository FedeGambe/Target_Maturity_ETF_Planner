"""Storico yield_etf.xlsx e foglio rendimenti_etf.xlsx."""
import csv
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from scripts.support import ALIQUOTA_CORP, ALIQUOTA_GOVT, BASE, PROVENTI, months_back, num, pulisci

HISTORY = BASE / "yield_etf.xlsx"
HEADER = ["Data estrazione", "Emittente", "Nome", "ISIN", "Tipo", "Proventi", "Scadenza",
          "TER", "Rendimento a scadenza", "Dati al", "Cedola"]  # Cedola aggiunta dopo: vuota nelle righe più vecchie
FORMATS = {"A": "DD/MM/YYYY", "G": "DD/MM/YYYY", "H": "0.00%", "I": "0.00%", "J": "DD/MM/YYYY", "K": "0.00%"}


def native(r):
    """Riga testuale (estrazione o CSV in sospeso: date gg/mm/aaaa, percentuali '2,75') -> tipi Excel."""
    d = lambda s: datetime.strptime(s, "%d/%m/%Y") if s else None
    f = lambda s: None if s in ("", None) else float(str(s).replace(",", ".")) / 100
    issuer = "iShares" if r[1].startswith("iShares") else r[1]
    # normalizza anche righe scritte con le etichette vecchie (CSV in sospeso)
    tipo_ = "Corp" if r[4] == "Corporate" else r[4]
    return [d(r[0]), issuer, pulisci(r[2]), r[3], tipo_, PROVENTI.get(r[5].lower(), r[5]),
            d(r[6]), f(r[7]), f(r[8]), d(r[9]), f(r[10]) if len(r) > 10 else None]


def append_history(rows):
    """Aggiunge allo storico. Se è aperto in Excel (bloccato) salva a parte in yield_etf_da_unire_*.csv:
    viene unito allo storico al giro successivo, così nessuna estrazione va persa."""
    pending = sorted(BASE.glob("yield_etf_da_unire_*.csv"))
    try:
        if HISTORY.exists():
            wb = load_workbook(HISTORY)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Storico"
            ws.append(HEADER)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col, width in zip("ABCDEFGHIJK", (15, 12, 70, 14, 10, 14, 11, 7, 20, 11, 9)):
                ws.column_dimensions[col].width = width
            ws.freeze_panes = "A2"
        first = ws.max_row + 1
        for p in pending:
            with open(p, encoding="utf-8-sig") as pf:
                for r in list(csv.reader(pf, delimiter=";"))[1:]:
                    ws.append(native(r))
        for r in rows:
            ws.append(native([*r[:7], num(r[7]), num(r[8]), r[9], num(r[10])]))
        for col, fmt in FORMATS.items():
            for cell in ws[col][first - 1:]:
                cell.number_format = fmt
        ws.auto_filter.ref = f"A1:K{ws.max_row}"
        wb.save(HISTORY)
    except PermissionError:
        p = BASE / f"yield_etf_da_unire_{datetime.now():%Y-%m-%d_%H%M}.csv"
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(HEADER)
            w.writerows([*r[:7], num(r[7]), num(r[8]), r[9], num(r[10])] for r in rows)
        print(f"{HISTORY.name} aperto in Excel: estrazione salvata in {p.name}, verrà unita al prossimo giro")
        return
    for p in pending:
        p.unlink()


def latest():
    """Ultimo dato disponibile di ogni ETF (per ISIN) dallo storico, con la variazione del rendimento lordo
    rispetto all'estrazione precedente, a un mese e a tre mesi prima (l'ultima estrazione a quella data o prima)."""
    by_isin = {}
    wb = load_workbook(HISTORY, read_only=True)
    for values in wb.active.iter_rows(min_row=2, values_only=True):
        r = dict(zip(HEADER, values))
        by_isin.setdefault(r["ISIN"], {})[r["Data estrazione"]] = r  # a parità di giorno vince la riga più recente
    wb.close()  # read_only tiene il file aperto finché non si chiude
    out = []
    for days in by_isin.values():
        dates = sorted(days)
        last = days[dates[-1]]

        def var(before):
            prev = [d for d in dates[:-1] if d <= before]
            if not prev or last["Rendimento a scadenza"] is None or days[prev[-1]]["Rendimento a scadenza"] is None:
                return None
            return round(last["Rendimento a scadenza"] - days[prev[-1]]["Rendimento a scadenza"], 6)

        last["Var 1g"] = var(dates[-1])
        last["Var 1m"] = var(months_back(dates[-1], 1))
        last["Var 3m"] = var(months_back(dates[-1], 3))
        out.append(last)
    return out


def build_sheet():
    """rendimenti_etf.xlsx dall'ultimo dato di ogni ETF nello storico. Se è aperto in Excel salva con l'ora nel nome."""
    path = BASE / "rendimenti_etf.xlsx"
    try:
        write_xlsx(path, latest())
    except PermissionError:
        path = path.with_name(f"rendimenti_etf_{datetime.now():%Y-%m-%d_%H%M}.xlsx")
        write_xlsx(path, latest())
    print(f"Foglio aggiornato: {path.name}")


def write_xlsx(path, rows):
    rows = [r for r in rows if "Rolling" not in r["Nome"]]  # i rolling non scadono davvero: restano solo nello storico

    def netto(r):  # stesso calcolo della colonna "Netto TER e tasse", con le aliquote di default
        aliquota = ALIQUOTA_GOVT if r["Tipo"] == "Govt" else ALIQUOTA_CORP
        return ((r["Rendimento a scadenza"] or 0) - (r["TER"] or 0)) * (1 - aliquota)

    rows.sort(key=netto, reverse=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rendimenti"
    # aliquote modificabili nel foglio (celle gialle): le colonne netto si ricalcolano da sole
    ws.append(["Aliquota Govt", ALIQUOTA_GOVT])
    ws.append(["Aliquota Corp", ALIQUOTA_CORP])
    ws.append([f"Aggiornato il {datetime.now():%d/%m/%Y %H:%M} da {HISTORY.name}"])
    # anagrafica | cedola, rendimento lordo e sue variazioni | costi e tasse | netti
    ws.append(["Emittente", "Nome", "ISIN", "Tipo", "Proventi", "Scadenza", "Anni a scadenza", "Cedola",
               "Rendimento lordo", "Var giorno prec", "Var mese prec", "Var trimestre prec",
               "TER", "Aliquota", "Rendimento netto", "Netto TER e tasse"])
    for i, r in enumerate(rows, start=5):
        ws.append([r["Emittente"], r["Nome"], r["ISIN"], r["Tipo"], r["Proventi"], r["Scadenza"],
                   f'=IF(F{i}="","",MAX(0,YEARFRAC(TODAY(),F{i})))',  # si aggiorna da solo ogni giorno
                   r.get("Cedola"), r["Rendimento a scadenza"], r["Var 1g"], r["Var 1m"], r["Var 3m"], r["TER"],
                   f'=IF(D{i}="Govt",$B$1,$B$2)',
                   f'=IF(I{i}="","",I{i}*(1-N{i}))',
                   f'=IF(I{i}="","",(I{i}-M{i})*(1-N{i}))'])  # il TER riduce il provento, le tasse si applicano dopo
    last = ws.max_row

    var = "+0.00%;-0.00%;0.00%"  # differenza in punti percentuali, col segno
    for col, fmt in (("F", "DD/MM/YYYY"), ("G", "0.0"), ("H", "0.00%"), ("I", "0.00%"), ("J", var), ("K", var),
                     ("L", var), ("M", "0.00%"), ("N", "0.0%"), ("O", "0.00%"), ("P", "0.00%")):
        for cell in ws[col][4:]:
            cell.number_format = fmt
    for col, width in zip("ABCDEFGHIJKLMNOP", (13, 78, 14, 7, 9, 11, 10, 8, 11, 11, 11, 13, 7, 9, 11, 12)):
        ws.column_dimensions[col].width = width

    # grafica: tabella Excel con righe alternate (porta anche i filtri), intestazione a capo
    ws.add_table(Table(displayName="Rendimenti", ref=f"A4:P{last}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))
    ws.row_dimensions[4].height = 32
    for cell in ws[4]:
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=5, min_col=4, max_col=16):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    for cell in ws["P"][4:]:
        cell.font = Font(bold=True)  # la colonna su cui è ordinato il foglio
    for c in ("A1", "A2"):
        ws[c].font = Font(bold=True)
    for c in ("B1", "B2"):
        ws[c].number_format = "0.0%"
        ws[c].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["A3"].font = Font(italic=True, color="808080")

    # variazioni: rosso se il rendimento scende, verde se sale, più intenso fino a ±0,25 punti
    for col in "JKL":
        ws.conditional_formatting.add(f"{col}5:{col}{last}", ColorScaleRule(
            start_type="num", start_value=-0.0025, start_color="F8696B",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=0.0025, end_color="63BE7B"))
    ws.freeze_panes = "C5"
    wb.save(path)
